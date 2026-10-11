"""Tests de operations/ocr.py — moteur embarqué, hors-ligne, accents français."""
import socket

import fitz
import pytest

pytest.importorskip("rapidocr")

from pdf_equilibrist.operations import ocr  # noqa: E402

PHRASE = "Société Équilibre : réglement déjà versé, façade où"


def _scanned_doc() -> fitz.Document:
    """PDF d'une page ne contenant qu'une image (aucune couche texte)."""
    src = fitz.open()
    sp = src.new_page()
    sp.insert_text((50, 100), PHRASE, fontsize=18)
    pix = sp.get_pixmap(matrix=fitz.Matrix(300 / 72, 300 / 72), alpha=False)
    doc = fitz.open()
    page = doc.new_page(width=sp.rect.width, height=sp.rect.height)
    page.insert_image(page.rect, pixmap=pix)
    return doc


@pytest.fixture
def no_network(monkeypatch):
    """Toute tentative de connexion ou de résolution DNS fait échouer le test."""
    def _refuse(*a, **k):
        raise AssertionError("connexion réseau tentée pendant l'OCR")
    monkeypatch.setattr(socket.socket, "connect", _refuse)
    monkeypatch.setattr(socket.socket, "connect_ex", _refuse)
    monkeypatch.setattr(socket, "getaddrinfo", _refuse)


def test_models_are_shipped_in_the_package():
    d = ocr._models_dir()
    for name in ocr._MODEL_FILES:
        assert (d / name).is_file(), name


def test_missing_models_raise_instead_of_downloading(monkeypatch, tmp_path):
    monkeypatch.setattr(ocr, "_models_dir", lambda: tmp_path)
    with pytest.raises(ocr.OcrModelsMissingError):
        ocr._get_ocr_engine()


def test_scanned_page_detection():
    doc = _scanned_doc()
    assert ocr._is_scanned_page(doc[0])


def test_apply_ocr_makes_page_searchable_offline(no_network):
    doc = _scanned_doc()
    assert doc[0].get_text().strip() == ""
    assert ocr.apply_ocr_to_pdf(doc) == 1
    text = doc[0].get_text()
    assert "Société" in text and "Équilibre" in text
    assert "déjà" in text and "façade" in text


def test_ocr_to_markdown_offline(no_network, tmp_path):
    out = ocr.ocr_to_markdown(_scanned_doc(), tmp_path / "out.md")
    assert "Société" in out.read_text(encoding="utf-8")


def test_text_pdf_is_not_ocred(no_network):
    doc = fitz.open()
    doc.new_page().insert_text((50, 100), "Déjà du texte", fontsize=14)
    assert ocr.apply_ocr_to_pdf(doc) == 0


def test_cli_selftest_passes_in_a_fresh_process(tmp_path):
    """`--ocr-selftest` (utilisé en CI sur l'exe figé) : processus isolé, car il coupe le réseau."""
    import subprocess
    import sys
    report = tmp_path / "report.txt"
    r = subprocess.run([sys.executable, "-m", "pdf_equilibrist.main", "--ocr-selftest", str(report)],
                       capture_output=True, text=True, timeout=300)
    assert r.returncode == 0, r.stderr
    assert report.read_text(encoding="utf-8").rstrip().endswith("OK")


def _docx_text(path) -> str:
    import docx
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        parts += [c.text for row in t.rows for c in row.cells]
    return "\n".join(parts)


def test_ocr_text_reaches_the_word_export(tmp_path):
    """Régression « OCR → Word » : le .docx doit contenir le texte reconnu, pas seulement l'image.

    Avant correction, to_word() rouvrait le PDF d'origine sur le disque et ignorait la couche
    OCR ajoutée en mémoire : un scan donnait un Word sans texte.
    """
    pytest.importorskip("pdf2docx")
    from pdf_equilibrist.operations import convert
    original = tmp_path / "scan.pdf"
    doc = _scanned_doc()
    doc.save(str(original))
    doc = fitz.open(str(original))          # comme l'application : document ouvert depuis le disque
    assert ocr.apply_ocr_to_pdf(doc) == 1   # couche OCR en mémoire uniquement
    out = tmp_path / "out.docx"
    convert.to_word(doc, original, out)
    text = _docx_text(out)
    assert "Société" in text and "Équilibre" in text, text


def test_word_export_includes_unsaved_in_memory_changes(tmp_path):
    pytest.importorskip("pdf2docx")
    from pdf_equilibrist.operations import convert
    original = tmp_path / "plain.pdf"
    d = fitz.open()
    d.new_page().insert_text((50, 100), "Texte enregistré", fontsize=14)
    d.save(str(original))
    doc = fitz.open(str(original))
    doc[0].insert_text((50, 140), "Ajout non enregistré", fontsize=14)
    out = tmp_path / "plain.docx"
    convert.to_word(doc, original, out)
    assert "Ajout non enregistré" in _docx_text(out)


def test_word_export_mixed_digital_and_scanned_pages(tmp_path):
    """Un PDF mêlant page numérique et page scannée : les deux textes sont là, le texte numérique une seule fois."""
    pytest.importorskip("pdf2docx")
    from pdf_equilibrist.operations import convert
    digital = fitz.open()
    digital.new_page().insert_text((50, 100), "Page numérique unique", fontsize=14)
    scanned = _scanned_doc()
    digital.insert_pdf(scanned)
    original = tmp_path / "mixed.pdf"
    digital.save(str(original))
    doc = fitz.open(str(original))
    assert ocr.apply_ocr_to_pdf(doc) == 1          # seule la page scannée est OCRisée
    out = tmp_path / "mixed.docx"
    convert.to_word(doc, original, out)
    text = _docx_text(out)
    assert text.count("Page numérique unique") == 1, text
    assert "Société" in text, text


def test_excel_export_rebuilds_a_scanned_table_from_ocr_positions(tmp_path):
    """Tableau scanné (traits dans l'image) : pdfplumber ne voit aucun trait, le repli par
    position du texte OCR doit tout de même reconstituer la grille."""
    openpyxl = pytest.importorskip("openpyxl")
    pytest.importorskip("pdfplumber")
    from pdf_equilibrist.operations import convert
    words = [["Alpha", "Bravo", "Charlie"], ["Delta", "Echo", "Foxtrot"], ["Golf", "Hotel", "India"]]
    src = fitz.open()
    sp = src.new_page()
    for r, row in enumerate(words):
        for c, word in enumerate(row):
            rect = fitz.Rect(50 + c * 150, 80 + r * 40, 200 + c * 150, 120 + r * 40)
            sp.draw_rect(rect, width=1)
            sp.insert_text((rect.x0 + 8, rect.y0 + 26), word, fontsize=16)
    pix = sp.get_pixmap(matrix=fitz.Matrix(300 / 72, 300 / 72), alpha=False)
    doc = fitz.open()
    page = doc.new_page(width=sp.rect.width, height=sp.rect.height)
    page.insert_image(page.rect, pixmap=pix)
    assert ocr.apply_ocr_to_pdf(doc) == 1
    out = tmp_path / "t.xlsx"
    convert.to_excel(doc, out)
    rows = [[c.value for c in r] for r in openpyxl.load_workbook(out).active.iter_rows()]
    assert rows == words, rows


def _two_scanned_pages() -> fitz.Document:
    doc = _scanned_doc()
    doc.insert_pdf(_scanned_doc())
    return doc


def test_cancel_leaves_the_document_strictly_untouched():
    """Annulation coopérative atomique : même si la 1re page est déjà reconnue, rien n'est inséré."""
    doc = _two_scanned_pages()
    calls = {"n": 0}

    def cancel_after_first_page():
        calls["n"] += 1
        return calls["n"] > 2          # 1er contrôle : page 1 ; 2e : page 2 ; 3e : annule

    with pytest.raises(ocr.OcrCancelled):
        ocr.apply_ocr_to_pdf(doc, cancel_cb=cancel_after_first_page)
    assert doc[0].get_text().strip() == "" and doc[1].get_text().strip() == ""
    assert not doc[0].get_texttrace() and not doc[1].get_texttrace()   # aucune couche OCR insérée


def test_cancel_before_start_does_nothing():
    doc = _scanned_doc()
    with pytest.raises(ocr.OcrCancelled):
        ocr.apply_ocr_to_pdf(doc, cancel_cb=lambda: True)
    assert doc[0].get_text().strip() == ""


def test_cancel_markdown_export_writes_no_file(tmp_path):
    out = tmp_path / "x.md"
    with pytest.raises(ocr.OcrCancelled):
        ocr.ocr_to_markdown(_scanned_doc(), out, cancel_cb=lambda: True)
    assert not out.exists()


def test_not_cancelled_ocr_still_processes_every_page():
    doc = _two_scanned_pages()
    assert ocr.apply_ocr_to_pdf(doc, cancel_cb=lambda: False) == 2
    assert "Société" in doc[0].get_text() and "Société" in doc[1].get_text()


# ── Tableaux scannés → Excel (grille détectée dans l'image) ───────────────────

_INVOICE = [
    ["Référence", "Désignation", "Qté", "Prix HT", "Total HT"],
    ["EQ-1001", "Bureau en chêne massif", "2", "420,00 €", "840,00 €"],
    ["EQ-1017", "Fauteuil ergonomique", "4", "189,50 €", "758,00 €"],
    ["EQ-2040", "Lampe de bureau LED", "6", "34,90 €", "209,40 €"],
]


def _invoice_doc() -> fitz.Document:
    """Page « scannée » : tableau à bordures + titre, uniquement sous forme d'image."""
    src = fitz.open()
    sp = src.new_page()
    sp.insert_text((50, 70), "FACTURE F-2026-0412", fontsize=18, fontname="hebo")
    widths = [80, 200, 40, 90, 90]
    for r, row in enumerate(_INVOICE):
        x = 50
        for c, cell in enumerate(row):
            rect = fitz.Rect(x, 110 + r * 28, x + widths[c], 110 + (r + 1) * 28)
            sp.draw_rect(rect, color=(0, 0, 0), width=1)
            sp.insert_text((rect.x0 + 6, rect.y0 + 18), cell, fontsize=11, fontname="helv")
            x += widths[c]
    pix = sp.get_pixmap(matrix=fitz.Matrix(200 / 72, 200 / 72), alpha=False)
    doc = fitz.open()
    page = doc.new_page(width=sp.rect.width, height=sp.rect.height)
    page.insert_image(page.rect, pixmap=pix)
    return doc


def _no_euro(rows):
    return [[c.replace(" €", "").replace("€", "").strip() for c in r] for r in rows]


def test_detect_grid_finds_the_ruled_lines():
    doc = _invoice_doc()
    xs, ys = ocr._detect_grid(ocr._page_to_pil(doc[0]))
    assert len(xs) == 6 and len(ys) == 5, (xs, ys)      # 5 colonnes / 4 lignes → 6 et 5 traits


def test_ocr_to_excel_rebuilds_a_ruled_scanned_table(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    pytest.importorskip("pdfplumber")
    doc = _invoice_doc()
    out = tmp_path / "facture.xlsx"
    ocr.ocr_to_excel(doc, out)
    rows = [[(c.value or "") for c in r] for r in openpyxl.load_workbook(out).active.iter_rows()]
    rows = [r for r in rows if any(r)]
    assert rows[0][0].startswith("FACTURE")                 # titre hors grille : ligne à une cellule
    assert _no_euro(rows[1:]) == _no_euro(_INVOICE), rows   # grille exacte, cellule par cellule
    # le symbole € instable du modèle est rétabli de façon homogène sur les colonnes de montants
    assert all(r[3].endswith("€") and r[4].endswith("€") for r in rows[2:]), rows
    # le document ouvert n'a pas été modifié (pas de couche OCR ajoutée)
    assert not doc[0].get_texttrace()


def test_ocr_to_excel_cancel_writes_nothing(tmp_path):
    out = tmp_path / "x.xlsx"
    with pytest.raises(ocr.OcrCancelled):
        ocr.ocr_to_excel(_invoice_doc(), out, cancel_cb=lambda: True)
    assert not out.exists()


def test_currency_repair_helpers():
    assert ocr._repair_currency_in_text("Total 1 284,60 · TTC") == "Total 1 284,60 € TTC"
    assert ocr._repair_currency_in_text("Version 1.2 · beta") == "Version 1.2 · beta"      # pas un montant
    assert ocr._repair_currency_in_cell("420,00.") == "420,00 €"
    assert ocr._repair_currency_in_cell("Fin de phrase.") == "Fin de phrase."
    table = [["Ref", "Prix"], ["A", "10,00 €"], ["B", "20,00"], ["C", "x"]]
    ocr._restore_currency_by_column(table)
    assert table == [["Ref", "Prix"], ["A", "10,00 €"], ["B", "20,00 €"], ["C", "x"]]
    nu = [["Qté"], ["2"], ["3,50"]]                                                        # aucune colonne en €
    ocr._restore_currency_by_column(nu)
    assert nu == [["Qté"], ["2"], ["3,50"]]
